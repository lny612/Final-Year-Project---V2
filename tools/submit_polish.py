"""Submission polish — final pass on Final_Report.docx.

Operations:
  1. Renumber body chapter / section / cross-reference numbering to match the
     Table of Contents the author edited:
         §6 → §4   (Project Background and Literature Review)
         §7 → §5   (Project Methodology)
         §8 → §6   (Results & Discussion)
         §9 → §7   (Conclusion and Future Works)
         §10 → §8  (Appendices)
         §11 → §9  (References)
     This applies to (a) the heading text itself ("7.3.6 Multi-model …" →
     "5.3.6 Multi-model …"), (b) every "§N(.M)*" / "Section N(.M)*" /
     "Table N(.M)*" reference inside body text and tables.
  2. Replace literal "\\pagebreak" text paragraphs with a real Word page break.
  3. Clean up the bracket figure placeholders:
       [Figure 7 — screenshot of the dossier panel …]   → "Figure 7. <new>"
       [Figure 4 — screenshots side-by-side of the three gate types …] →
                                                              "Figure 4. <new>"
       Figure 8 — sample wand stylized-illustration outputs …  →
                                                              "Figure 8. <new>"
  4. Polish the in-game screenshot captions Figure 2.1–2.8.

Backup is written to Final_Report.docx.bak6 before any edit.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH    = PROJECT_ROOT / ".claude" / "Final Report" / "Final_Report.docx"
BACKUP_PATH  = PROJECT_ROOT / ".claude" / "Final Report" / "Final_Report.docx.bak6"

BODY_FONT = "Times New Roman"

# Chapter-number remap (used for §, Section, Table, and heading text)
CHAPTER_MAP = {6: 4, 7: 5, 8: 6, 9: 7, 10: 8, 11: 9}


def set_run_font(run, *, bold=False, italic=False, size=None):
    run.font.name = BODY_FONT
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), BODY_FONT)
    if bold:   run.bold = True
    if italic: run.italic = True
    if size:   run.font.size = size


def add_run(p, text, **kw):
    r = p.add_run(text)
    set_run_font(r, **kw)
    return r


def replace_runs(p, runs):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, kw in runs:
        add_run(p, text, **kw)


# ---------------------------------------------------------------------------
# Renumbering
# ---------------------------------------------------------------------------
# Regex matches:  §N, §N.M, §N.M.K   |   Section N, Section N.M, Section N.M.K
# |   Table N.M(.K)?
SECTION_REF_RX = re.compile(
    r"(?P<lead>§|Section\s+|Table\s+)(?P<num>\d+)(?P<rest>(?:\.\d+){0,2})"
)
# Heading text that starts with the chapter number, e.g. "6. Project …",
# "6.3 Diffusion …", "6.3.1 LLM …".
HEADING_RX = re.compile(
    r"^(?P<num>\d+)(?P<rest>(?:\.\d+){0,2})(?P<sep>(?:\.\s)|(?:\s))"
)


def remap_section_refs_in_text(text: str) -> tuple[str, int]:
    """Return (new_text, replacement_count)."""
    n = 0

    def repl(m):
        nonlocal n
        num = int(m.group("num"))
        if num not in CHAPTER_MAP:
            return m.group(0)
        n += 1
        return f"{m.group('lead')}{CHAPTER_MAP[num]}{m.group('rest')}"

    return SECTION_REF_RX.sub(repl, text), n


def remap_heading_text(text: str) -> tuple[str, int]:
    """If the text starts with a chapter number we should remap, do so."""
    m = HEADING_RX.match(text)
    if not m:
        return text, 0
    num = int(m.group("num"))
    if num not in CHAPTER_MAP:
        return text, 0
    new_num = CHAPTER_MAP[num]
    new_prefix = f"{new_num}{m.group('rest')}{m.group('sep')}"
    return new_prefix + text[m.end():], 1


def renumber_paragraph(p, *, is_heading: bool):
    """Renumber section refs (and the heading-prefix number when applicable)
    in-place, preserving run-level formatting where possible."""
    full = "".join(r.text for r in p.runs)

    new_full = full
    n = 0
    if is_heading:
        new_full, k = remap_heading_text(new_full)
        n += k
    new_full, k = remap_section_refs_in_text(new_full)
    n += k
    if not n:
        return 0

    # Try per-run replacement so formatting (bold/italic) is preserved when the
    # match lives entirely inside one run.
    if len(p.runs) == 1:
        p.runs[0].text = new_full
        return n
    # Multi-run: walk runs and apply the same regexes per run (this catches
    # most refs since they rarely span runs).
    for r in p.runs:
        if not r.text:
            continue
        if is_heading and r is p.runs[0]:
            r.text, _ = remap_heading_text(r.text)
        r.text, _ = remap_section_refs_in_text(r.text)

    rebuilt = "".join(r.text for r in p.runs)
    if rebuilt != new_full:
        # Fallback: collapse to a single run
        replace_runs(p, [(new_full, {})])
    return n


def renumber_table_cell(cell):
    n = 0
    for p in cell.paragraphs:
        n += renumber_paragraph(p, is_heading=False)
    return n


# ---------------------------------------------------------------------------
def replace_pagebreak_paragraph(p):
    """Replace the literal '\\pagebreak' text with a real Word page break run."""
    text = "".join(r.text for r in p.runs)
    if "\\pagebreak" not in text:
        return False
    # Wipe runs and add a single empty run carrying the page break.
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run("")
    run.add_break(WD_BREAK.PAGE)
    return True


# ---------------------------------------------------------------------------
def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not DOCX_PATH.exists():
        sys.exit(f"missing docx: {DOCX_PATH}")
    print(f"backup -> {BACKUP_PATH}")
    shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(str(DOCX_PATH))
    print(f"opened {DOCX_PATH} ({len(doc.paragraphs)} paragraphs, "
          f"{len(doc.tables)} tables)")

    # -----------------------------------------------------------------------
    # 1. Renumber every paragraph + every table cell.
    # -----------------------------------------------------------------------
    n_total = 0
    for p in doc.paragraphs:
        is_h = p.style.name.startswith("Heading")
        n_total += renumber_paragraph(p, is_heading=is_h)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                n_total += renumber_table_cell(cell)
    print(f"  [renumber] applied {n_total} substitutions")

    # -----------------------------------------------------------------------
    # 2. Replace the two literal \pagebreak paragraphs.
    # -----------------------------------------------------------------------
    n_pb = 0
    for p in doc.paragraphs:
        if replace_pagebreak_paragraph(p):
            n_pb += 1
    print(f"  [pagebreak] converted {n_pb} literal paragraphs to page breaks")

    # -----------------------------------------------------------------------
    # 3. Clean bracket figure placeholders.
    # -----------------------------------------------------------------------
    # Figure 7 — dossier panel screenshot (now §5.5.1 — was §7.5.1)
    for p in doc.paragraphs:
        if p.text.startswith("[Figure 7"):
            print(f"  [fig7] replacing dossier-panel placeholder")
            replace_runs(p, [
                ("Figure 7. Memo-gated dossier reading. The player drags across "
                 "phrases in the customer's prose to highlight them, then drops "
                 "each highlight into one of four memo slots (Element / "
                 "Personality / Purpose / Reinforcement). Proceed enables once "
                 "every slot has at least one chip.", {"italic": True, "size": Pt(10)}),
            ])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # Figure 4 — three-gate screenshot
    for p in doc.paragraphs:
        if p.text.startswith("[Figure 4"):
            print(f"  [fig4] replacing tracing-minigame placeholder")
            replace_runs(p, [
                ("Figure 4. Three gate types in the tracing minigame. "
                 "Left: amber-square Tap with a key letter — single-press to "
                 "clear. Centre: cyan-square Hold with a green halo and partial "
                 "fill — sustained press for 0.9 s. Right: magenta-square "
                 "Accent with a directional chevron and arrow glyph — press "
                 "the key, then flick the cursor in the indicated compass "
                 "direction.", {"italic": True, "size": Pt(10)}),
            ])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # Figure 8 — sample wand outputs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Figure 8 —") or t.startswith("[Figure 8"):
            print(f"  [fig8] polishing wand-output caption")
            replace_runs(p, [
                ("Figure 8. Sample wand stylized-illustrations from the live "
                 "ComfyUI runtime. Four 128×128 outputs across customer schools "
                 "(Storm, Hydromancy, Shadow, Pyromancy) showing the "
                 "visual diversity the SVDQuant-quantised Z-Image Turbo model "
                 "produces in 3–6 s on the project's 8 GB GPU.",
                 {"italic": True, "size": Pt(10)}),
            ])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    # -----------------------------------------------------------------------
    # 4. Polish Figure 2.1 – 2.8 screenshot captions (now under §6.1).
    # -----------------------------------------------------------------------
    SCREENSHOT_CAPTIONS = {
        "Figure 2.1 Title":
            ("Figure 2.1. Title scene — parchment menu with Start and Quit "
             "buttons; entry point to a new seven-day session."),
        "Figure 2.2 Morning":
            ("Figure 2.2. Morning scene — letter typewriter with sender wax "
             "seal; customer and material pre-generation cascade silently in "
             "the background while the player reads."),
        "Figure 2.3 Customer (dossier)":
            ("Figure 2.3. Customer scene — drag-to-highlight memo gate. "
             "Phrases pulled from the dossier (left) snap into one of four "
             "memo slots (right); Proceed enables once every slot has at least "
             "one chip."),
        "Figure 2.4 Market":
            ("Figure 2.4. Market scene — six AI-generated materials (three "
             "cores + three woods). ✦ glyphs and inline blue-tint mark cards "
             "whose description fields token-match the player's memo."),
        "Figure 2.5 Crafting":
            ("Figure 2.5. Crafting scene — two-core plus one-wood slot "
             "layout. The memo card is shown alongside the slots, replacing "
             "the dossier as the in-scene reference."),
        "Figure 2.6 Minigame mid-round":
            ("Figure 2.6. Tracing minigame, round 3 in progress. Red mist "
             "chases the blue player-traced fill; cyan Hold and magenta "
             "Accent gates are visible at conductor-ictus points along the "
             "rune path."),
        "Figure 2.7 Evaluation Reveal":
            ("Figure 2.7. Evaluation scene — theatrical reveal of the "
             "AI-generated wand, scoreboard (Conjuring / Materials / Customer "
             "Fit / Reward), and final letter grade."),
        "Figure 2.8 Royal Ending":
            ("Figure 2.8. Royal ending — pre-generated illustration plus "
             "typewritten dialogue, triggered when peak reputation reaches "
             "the Royal threshold (≥ 90)."),
    }

    n_caps = 0
    for p in doc.paragraphs:
        key = p.text.strip()
        if key in SCREENSHOT_CAPTIONS:
            v = SCREENSHOT_CAPTIONS[key]
            replace_runs(p, [(v, {"italic": True, "size": Pt(10)})])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p.style = doc.styles["Normal"]
            except KeyError:
                pass
            n_caps += 1
            print(f"  [fig2.x] polished: {key}")

    print(f"  [fig2.x] polished {n_caps} screenshot captions")

    # -----------------------------------------------------------------------
    # 5. Sanity audit of remaining literal "\\pagebreak" or "[Figure" text.
    # -----------------------------------------------------------------------
    leftovers = []
    for i, p in enumerate(doc.paragraphs):
        if "\\pagebreak" in p.text:
            leftovers.append((i, "pagebreak", p.text[:80]))
        if re.match(r"\[Figure \d", p.text):
            leftovers.append((i, "bracket-figure", p.text[:80]))
    if leftovers:
        print("  [audit] leftovers found:")
        for i, kind, t in leftovers:
            print(f"    {i:4d} [{kind}] {t}")
    else:
        print("  [audit] no leftover placeholders or pagebreak literals")

    print(f"saving -> {DOCX_PATH}")
    doc.save(str(DOCX_PATH))
    print("done")


if __name__ == "__main__":
    sys.exit(main() or 0)
